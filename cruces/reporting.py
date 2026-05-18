from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .analysis_engine import build_executive_analysis


def _value(value: Any, default: str = ""):
    return default if value is None else value


def _autosize_sheet(sheet):
    for column_cells in sheet.columns:
        max_length = 0
        column = get_column_letter(column_cells[0].column)
        for cell in column_cells:
            max_length = max(max_length, len(str(cell.value or "")))
        sheet.column_dimensions[column].width = min(max(max_length + 2, 12), 45)


def build_analysis_excel(analysis) -> bytes:
    summary = analysis.summary or {}
    rows = analysis.route_rows or []
    executive = summary.get("executive_analysis") or build_executive_analysis(summary, rows)

    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Resumen"
    summary_sheet.append(["Campo", "Valor"])
    summary_items = [
        ("Análisis", analysis.name),
        ("Cruce", analysis.cruce.interseccion),
        ("Código J", analysis.cruce.codigo_j),
        ("Resultado", analysis.get_result_display()),
        ("Riesgo operacional", executive.get("risk", {}).get("label")),
        ("Horario", f"{analysis.horario_inicio:%H:%M} - {analysis.horario_fin:%H:%M}"),
        ("Variación velocidad (%)", summary.get("speed_delta_pct")),
        ("Variación tiempo (%)", summary.get("time_delta_pct")),
        ("Velocidad antes (km/h)", summary.get("avg_speed_before")),
        ("Velocidad después (km/h)", summary.get("avg_speed_after")),
        ("Rutas analizadas", summary.get("routes_analyzed")),
        ("Resumen ejecutivo", executive.get("headline")),
        ("Conclusión", summary.get("conclusion")),
    ]
    for item in summary_items:
        summary_sheet.append(item)
    for cell in summary_sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1D4ED8")
    for row in range(1, summary_sheet.max_row + 1):
        summary_sheet[f"B{row}"].alignment = Alignment(wrap_text=True, vertical="top")
    _autosize_sheet(summary_sheet)

    executive_sheet = workbook.create_sheet("Analisis ejecutivo")
    executive_sheet.append(["Seccion", "Detalle"])
    executive_sheet.append(["Resumen", executive.get("headline")])
    executive_sheet.append(["Narrativa", executive.get("narrative")])
    executive_sheet.append(["Nivel de riesgo", executive.get("risk", {}).get("label")])
    executive_sheet.append(["Motivo riesgo", executive.get("risk", {}).get("reason")])
    executive_sheet.append(["Calidad datos", executive.get("quality", {}).get("label")])
    executive_sheet.append(["Lecturas clave", ""])
    for item in executive.get("readouts", []):
        executive_sheet.append(["", item])
    executive_sheet.append(["Recomendaciones", ""])
    for item in executive.get("recommendations", []):
        executive_sheet.append(["", item])
    executive_sheet.append(["Rutas criticas", "Velocidad %, Tiempo %, Muestras, Confianza"])
    for row in executive.get("critical_routes", []):
        executive_sheet.append(
            [
                row.get("route"),
                f"{row.get('speed_delta_pct')}%, {row.get('time_delta_pct')}%, "
                f"{row.get('samples_before')}/{row.get('samples_after')}, {row.get('confidence_label')}",
            ]
        )
    executive_sheet.append(["Mejores mejoras", "Velocidad %, Tiempo %, Muestras"])
    for row in executive.get("top_improvements", []):
        executive_sheet.append(
            [
                row.get("route"),
                f"{row.get('speed_delta_pct')}%, {row.get('time_delta_pct')}%, "
                f"{row.get('samples_before')}/{row.get('samples_after')}",
            ]
        )
    for cell in executive_sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1D4ED8")
    for row in range(1, executive_sheet.max_row + 1):
        executive_sheet[f"B{row}"].alignment = Alignment(wrap_text=True, vertical="top")
    _autosize_sheet(executive_sheet)

    detail_sheet = workbook.create_sheet("Detalle rutas")
    detail_sheet.append(
        [
            "Ruta",
            "Obs. antes",
            "Obs. después",
            "Vel. antes",
            "Vel. después",
            "Delta vel. %",
            "Tiempo antes",
            "Tiempo después",
            "Delta tiempo %",
            "Resultado",
            "P15 antes",
            "P85 antes",
            "P15 después",
            "P85 después",
        ]
    )
    for row in rows:
        detail_sheet.append(
            [
                row.get("route"),
                row.get("samples_before"),
                row.get("samples_after"),
                row.get("speed_before"),
                row.get("speed_after"),
                row.get("speed_delta_pct"),
                row.get("time_before"),
                row.get("time_after"),
                row.get("time_delta_pct"),
                row.get("result_label"),
                row.get("speed_p15_before"),
                row.get("speed_p85_before"),
                row.get("speed_p15_after"),
                row.get("speed_p85_after"),
            ]
        )
    for cell in detail_sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="334155")
    _autosize_sheet(detail_sheet)

    quality_sheet = workbook.create_sheet("Calidad")
    quality_sheet.append(["Archivo", "Filas totales", "Procesadas", "Fuera horario", "Inválidas", "Rutas"])
    for label, key in (("ANTES", "before_quality"), ("DESPUÉS", "after_quality")):
        quality = summary.get(key, {})
        quality_sheet.append(
            [
                label,
                quality.get("total_rows"),
                quality.get("processed_rows"),
                quality.get("out_of_window_rows"),
                quality.get("invalid_rows"),
                quality.get("routes"),
            ]
        )
    for cell in quality_sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="334155")
    _autosize_sheet(quality_sheet)

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def build_analysis_pdf(analysis) -> bytes:
    summary = analysis.summary or {}
    rows = analysis.route_rows or []
    executive = summary.get("executive_analysis") or build_executive_analysis(summary, rows)
    output = BytesIO()
    doc = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        rightMargin=1.2 * cm,
        leftMargin=1.2 * cm,
        topMargin=1.2 * cm,
        bottomMargin=1.2 * cm,
    )
    styles = getSampleStyleSheet()
    story = [
        Paragraph(f"Reporte de análisis: {analysis.name}", styles["Title"]),
        Paragraph(f"Cruce: {analysis.cruce.interseccion}", styles["Heading2"]),
        Paragraph(_value(summary.get("conclusion"), "Sin conclusión disponible."), styles["BodyText"]),
        Spacer(1, 0.35 * cm),
    ]

    summary_table = Table(
        [
            ["Resultado", analysis.get_result_display(), "Horario", f"{analysis.horario_inicio:%H:%M} - {analysis.horario_fin:%H:%M}"],
            ["Delta velocidad", f"{summary.get('speed_delta_pct', 0)}%", "Delta tiempo", f"{summary.get('time_delta_pct', 0)}%"],
            ["Vel. antes", f"{summary.get('avg_speed_before', 0)} km/h", "Vel. después", f"{summary.get('avg_speed_after', 0)} km/h"],
            ["Rutas analizadas", summary.get("routes_analyzed", 0), "Rutas comunes", summary.get("common_routes", 0)],
        ],
        colWidths=[4 * cm, 4 * cm, 4 * cm, 4 * cm],
    )
    summary_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EEF8")),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
                ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
                ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
                ("FONTNAME", (2, 0), (2, -1), "Helvetica-Bold"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("PADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    story.extend([summary_table, Spacer(1, 0.45 * cm)])

    story.extend(
        [
            Paragraph("Resumen ejecutivo", styles["Heading2"]),
            Paragraph(_value(executive.get("headline"), "Sin resumen ejecutivo disponible."), styles["BodyText"]),
            Paragraph(_value(executive.get("narrative"), ""), styles["BodyText"]),
            Paragraph(
                f"Nivel de riesgo: {executive.get('risk', {}).get('label', 'Sin datos')}. "
                f"{executive.get('risk', {}).get('reason', '')}",
                styles["BodyText"],
            ),
            Spacer(1, 0.25 * cm),
            Paragraph("Lecturas clave", styles["Heading3"]),
        ]
    )
    for item in executive.get("readouts", []):
        story.append(Paragraph(f"- {item}", styles["BodyText"]))
    story.append(Spacer(1, 0.2 * cm))
    story.append(Paragraph("Recomendaciones", styles["Heading3"]))
    for item in executive.get("recommendations", []):
        story.append(Paragraph(f"- {item}", styles["BodyText"]))
    story.append(Spacer(1, 0.35 * cm))

    if executive.get("critical_routes"):
        critical_table_data = [["Ruta critica", "Delta vel.", "Delta tiempo", "Muestras", "Confianza"]]
        for row in executive.get("critical_routes", [])[:5]:
            critical_table_data.append(
                [
                    str(row.get("route", ""))[:60],
                    f"{row.get('speed_delta_pct', 0)}%",
                    f"{row.get('time_delta_pct', 0)}%",
                    f"{row.get('samples_before', 0)}/{row.get('samples_after', 0)}",
                    row.get("confidence_label", ""),
                ]
            )
        critical_table = Table(critical_table_data, repeatRows=1)
        critical_table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#7F1D1D")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#CBD5E1")),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, -1), 7),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("PADDING", (0, 0), (-1, -1), 4),
                ]
            )
        )
        story.extend([Paragraph("Rutas priorizadas", styles["Heading3"]), critical_table, Spacer(1, 0.35 * cm)])

    route_table_data = [["Ruta", "Obs.", "Vel. antes", "Vel. después", "Delta vel.", "Tiempo antes", "Tiempo después", "Resultado"]]
    for row in rows[:35]:
        route_table_data.append(
            [
                str(row.get("route", ""))[:52],
                f"{row.get('samples_before', 0)}/{row.get('samples_after', 0)}",
                row.get("speed_before", 0),
                row.get("speed_after", 0),
                f"{row.get('speed_delta_pct', 0)}%",
                row.get("time_before", 0),
                row.get("time_after", 0),
                row.get("result_label", ""),
            ]
        )
    route_table = Table(route_table_data, repeatRows=1)
    route_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#334155")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#CBD5E1")),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 7),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("PADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    story.extend([Paragraph("Detalle de rutas", styles["Heading2"]), route_table])
    doc.build(story)
    return output.getvalue()
