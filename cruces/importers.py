from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel


HEADER_ALIASES = {
    "fecha": "fecha_terreno",
    "jornada": "jornada",
    "horarioinicio": "horario_inicio",
    "horariofin": "horario_fin",
    "solicita": "solicita",
    "proyectoobra": "project_name",
    "proyecto": "project_name",
    "obra": "project_name",
    "codigo": "codigo_j",
    "codigoj": "codigo_j",
    "codigocruce": "codigo_j",
    "interseccion": "interseccion",
    "cruce": "interseccion",
    "comuna": "comuna",
    "evaluador": "evaluador",
    "estado": "estado",
    "estadodesemaforo": "estado_semaforo",
    "estadosemaforo": "estado_semaforo",
    "serealizomodificacion": "realizo_modificacion",
    "serealizomodif": "realizo_modificacion",
    "modificacion": "modificacion_texto",
    "observaciones": "observaciones",
    "observacion": "observaciones",
    "poligono": "poligono",
    "rutas": "rutas_texto",
}

REQUIRED_HEADER_KEYS = {"fecha_terreno", "interseccion"}


@dataclass(frozen=True)
class BitacoraParseResult:
    sheet_name: str
    header_row: int
    headers: list[str]
    rows: list[dict[str, Any]]


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return re.sub(r"[^a-zA-Z0-9]+", "", text).lower()


def compact_string(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.strftime("%H:%M")
    return re.sub(r"\s+", " ", str(value).strip())


def clean_placeholder(value: Any, placeholders: set[str]) -> str:
    text = compact_string(value)
    if normalize_text(text) in placeholders:
        return ""
    return text


def parse_date_value(value: Any) -> str:
    if value in (None, ""):
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)):
        try:
            return from_excel(value).date().isoformat()
        except (TypeError, ValueError):
            return ""

    text = compact_string(value)
    if not text:
        return ""

    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            continue
    return ""


def _time_from_match(text: str) -> str:
    match = re.search(r"(\d{1,2})\s*[:.]\s*(\d{2})", text)
    if not match:
        return ""
    hour = int(match.group(1))
    minute = int(match.group(2))
    if hour > 23 or minute > 59:
        return ""
    return f"{hour:02d}:{minute:02d}"


def parse_time_value(value: Any) -> str:
    if value in (None, ""):
        return ""
    if isinstance(value, datetime):
        return value.strftime("%H:%M")
    if isinstance(value, time):
        return value.strftime("%H:%M")
    if isinstance(value, (int, float)):
        if 0 <= float(value) < 1:
            minutes = round(float(value) * 24 * 60)
            return (datetime(2000, 1, 1) + timedelta(minutes=minutes)).strftime("%H:%M")
        return ""
    return _time_from_match(compact_string(value))


def parse_time_range(start_value: Any, end_value: Any) -> tuple[str, str]:
    start_text = compact_string(start_value)
    if start_text and re.search(r"\d{1,2}\s*[:.]\s*\d{2}\s*[-–—]\s*\d{1,2}\s*[:.]\s*\d{2}", start_text):
        parts = re.split(r"[-–—]", start_text, maxsplit=1)
        return parse_time_value(parts[0]), parse_time_value(parts[1])
    return parse_time_value(start_value), parse_time_value(end_value)


def parse_bool_value(value: Any) -> bool:
    text = normalize_text(value)
    return text in {"si", "s", "yes", "true", "1", "x"}


def normalize_jornada(value: Any) -> str:
    text = normalize_text(value)
    if text.startswith("pmd"):
        return "PMD"
    if text.startswith("pm"):
        return "PM"
    if text.startswith("pt"):
        return "PT"
    if text.startswith("fp"):
        return "FP"
    return compact_string(value).upper()


def normalize_estado_semaforo(value: Any) -> str:
    text = normalize_text(value)
    if "aislado" in text:
        return "Aislado"
    if "sistema" in text:
        return "En Sistema"
    return compact_string(value)


def normalize_estado(value: Any) -> str:
    text = normalize_text(value)
    if "realizado" in text:
        return "Realizado"
    if "programado" in text:
        return "Programado"
    if "extraprogramatico" in text or "evento" in text:
        return "Evento Extraprogramático"
    if "confirmar" in text:
        return "Por confirmar"
    return compact_string(value)


def _build_header_map(values: list[Any]) -> dict[int, str]:
    mapping = {}
    used_fields = set()
    for index, value in enumerate(values):
        normalized = normalize_text(value)
        field_name = HEADER_ALIASES.get(normalized)
        if field_name and field_name not in used_fields:
            mapping[index] = field_name
            used_fields.add(field_name)
    return mapping


def _find_header_sheet(workbook) -> tuple[Any, int, dict[int, str], list[str]]:
    for worksheet in workbook.worksheets:
        for row_index, row in enumerate(worksheet.iter_rows(max_row=40, values_only=True), start=1):
            values = list(row)
            header_map = _build_header_map(values)
            if REQUIRED_HEADER_KEYS.issubset(set(header_map.values())):
                headers = [compact_string(value) for value in values if compact_string(value)]
                return worksheet, row_index, header_map, headers
    raise ValueError("No se encontró una hoja con encabezados de bitácora válidos.")


def parse_bitacora(path: str | Path) -> BitacoraParseResult:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet, header_row, header_map, headers = _find_header_sheet(workbook)
        parsed_rows: list[dict[str, Any]] = []

        for row_number, row in enumerate(
            worksheet.iter_rows(min_row=header_row + 1, values_only=True),
            start=header_row + 1,
        ):
            raw = {field_name: row[index] if index < len(row) else None for index, field_name in header_map.items()}
            interseccion = compact_string(raw.get("interseccion"))
            codigo_j = compact_string(raw.get("codigo_j")).upper()
            if not interseccion and not codigo_j:
                continue

            horario_inicio, horario_fin = parse_time_range(raw.get("horario_inicio"), raw.get("horario_fin"))
            warnings = []
            if not interseccion:
                warnings.append("Sin intersección.")
            if not codigo_j:
                warnings.append("Sin código J.")
            if not parse_date_value(raw.get("fecha_terreno")):
                warnings.append("Sin fecha válida.")
            rutas_texto = clean_placeholder(raw.get("rutas_texto"), {"rutas", "ruta"})
            if rutas_texto:
                warnings.append("Trae rutas asociadas.")

            parsed_rows.append(
                {
                    "row_number": row_number,
                    "codigo_j": codigo_j,
                    "interseccion": interseccion or codigo_j,
                    "comuna": compact_string(raw.get("comuna")),
                    "fecha_terreno": parse_date_value(raw.get("fecha_terreno")),
                    "evaluador": compact_string(raw.get("evaluador")),
                    "jornada": normalize_jornada(raw.get("jornada")),
                    "horario_inicio": horario_inicio,
                    "horario_fin": horario_fin,
                    "solicita": compact_string(raw.get("solicita")),
                    "project_name": compact_string(raw.get("project_name")),
                    "estado": normalize_estado(raw.get("estado")),
                    "estado_semaforo": normalize_estado_semaforo(raw.get("estado_semaforo")),
                    "realizo_modificacion": parse_bool_value(raw.get("realizo_modificacion")),
                    "modificacion_texto": compact_string(raw.get("modificacion_texto")),
                    "observaciones": compact_string(raw.get("observaciones")),
                    "rutas_texto": rutas_texto,
                    "warnings": warnings,
                    "status": "ready",
                    "status_label": "Listo",
                    "status_detail": "",
                }
            )

        return BitacoraParseResult(
            sheet_name=worksheet.title,
            header_row=header_row,
            headers=headers,
            rows=parsed_rows,
        )
    finally:
        workbook.close()
